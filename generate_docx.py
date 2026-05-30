import docx
import os
import re

def add_markdown_to_doc(doc, md_text):
    for line in md_text.split('\n'):
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.strip() != "":
            # Basic bold parsing for simple paragraphs
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith('`') and part.endswith('`'):
                    p.add_run(part).font.name = 'Courier New'
                else:
                    p.add_run(part)

doc = docx.Document()
doc.add_heading('Actualización de Documentación: Relojería IA v2.1', 0)

files_to_read = [
    ("README.md", "README.md"),
    ("Requerimientos (PRD)", "docs/PRD.md"),
    ("Matriz de Trazabilidad (RTM)", "docs/RTM.md"),
    ("Diccionario de Datos", "docs/DATABASE.md"),
    ("Arquitectura", "docs/ARCHITECTURE.md"),
    ("Plan de Pruebas QA (UI/UX)", "docs/QA_UI_TEST_PLAN.md"),
    ("Formato Pruebas UX (Ficticio)", "docs/Formato_Pruebas_UX_Datos_Ficticios.md")
]

for title, filepath in files_to_read:
    doc.add_heading(title, level=1)
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            add_markdown_to_doc(doc, f.read())
    except Exception as e:
        doc.add_paragraph(f"Error leyendo {filepath}: {e}")
    doc.add_page_break()

doc.save('Anexo_Documentacion_v2.docx')
print("Archivo DOCX generado con éxito.")
